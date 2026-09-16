from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl


root = Path(sys.argv[1]).resolve()
forecast = root / "Nigeria Quarterly GDP Nowcast.xlsx"
results = json.loads((root / ".automation" / "nowcast_results.json").read_text(encoding="utf-8"))
master = Path(results["source"])

assert forecast.exists() and forecast.stat().st_size > 50_000, forecast
assert master.exists() and master.stat().st_size > 30_000, master

wb = openpyxl.load_workbook(forecast, data_only=True, read_only=True)
assert "Executive Summary" in wb.sheetnames
assert "Quarterly Forecast" in wb.sheetnames
assert "GDP Interpolation" not in wb.sheetnames
monthly_ws = wb["Monthly Inputs"]
assert monthly_ws["K1"].value == "Inflation risk"
assert monthly_ws["L1"].value == "PMI"
monthly_rows = list(monthly_ws.iter_rows(min_row=2, values_only=True))
for row in monthly_rows:
    if row[0] is not None and row[0].year < 2015:
        assert row[10] is None, f"Inflation risk must be blank before January 2015: {row[0]}"
jan_2015 = next(row for row in monthly_rows if row[0] is not None and row[0].year == 2015 and row[0].month == 1)
assert abs(jan_2015[10] - 5.22190321364737) < 1e-8
assert abs(jan_2015[11] - 50.225124280907416) < 1e-8
internal_point = wb["Model Engine"]["H11"].value
release_mode = wb["Model Engine"]["H20"].value
complete_months = wb["Model Engine"]["H18"].value
pmi = wb["Executive Summary"]["C7"].value
forecast_ws = wb["Quarterly Forecast"]
forecast_record = next(
    row for row in forecast_ws.iter_rows(min_row=4, values_only=True)
    if row[0] and row[1] is None and str(row[8] or "").lower() != "official actual"
)
published_point = forecast_record[5]
published_lower = forecast_record[6]
published_upper = forecast_record[7]
assert isinstance(internal_point, (int, float))
assert isinstance(pmi, (int, float))
assert 0.0 < internal_point < 15.0
assert 0.0 < pmi < 100.0
assert isinstance(complete_months, (int, float)) and 0 <= complete_months <= 3
if complete_months < 3:
    assert release_mode == "Indicative range"
    assert published_point is None
    assert abs(published_lower - 4.34) < 1e-8
    assert abs(published_upper - 4.50) < 1e-8
else:
    assert release_mode == "Point forecast"
    assert isinstance(published_point, (int, float))
wb.close()

wb = openpyxl.load_workbook(master, data_only=False, read_only=True)
assert wb["monthly data"]["K2"].value == "PMI"
assert wb["monthly data"]["L2"].value == "Inflation Risk"
assert "Quarterly Data" in wb.sheetnames
wb.close()

print({"status": "refresh artifacts verified", "release_mode": release_mode, "internal_point": internal_point, "published_range": [published_lower, published_upper], "complete_months": complete_months, "pmi": pmi})
