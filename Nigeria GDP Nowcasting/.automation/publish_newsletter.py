from __future__ import annotations

import argparse
import copy
import json
import math
import re
from datetime import datetime
from pathlib import Path

import openpyxl
from docx import Document
from PIL import Image, ImageDraw, ImageFont


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_TEMPLATE = (
    PROJECT_ROOT
    / "CAPE ERC Economic Newsletter"
    / "2026"
    / "September Newsletter"
    / "CAPE Economic Performance and Prospect Bulletin September 2026.docx"
)
DEFAULT_OUTPUT = DEFAULT_TEMPLATE


def quarter_label(value: str) -> str:
    value = str(value).strip().upper()
    if "Q" in value:
        year, quarter = value.split("Q", 1)
        return f"Q{quarter} {year}"
    return value


def read_nowcast(workbook_path: Path) -> dict:
    wb = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    summary = wb["Executive Summary"]
    forecast = wb["Quarterly Forecast"]
    engine = wb["Model Engine"]
    monthly = wb["Monthly Inputs"]
    features = []
    for row in engine.iter_rows(min_row=5, max_col=1, values_only=True):
        if row[0] is None:
            break
        features.append(str(row[0]))
    feature_labels = {
        "gdp_lag1": "lagged GDP growth",
        "gdp_lag4": "year-earlier GDP growth",
        "headline_yoy": "headline inflation",
        "core_yoy": "core inflation",
        "food_yoy": "food inflation",
        "cpd": "crude oil production",
        "pmi": "PMI",
        "risk": "inflation risk",
        "m2": "M2",
        "m2_yoy": "annual M2 growth",
        "m2_mom": "monthly M2 growth",
        "m2_qoq": "quarterly M2 growth",
    }
    readable_features = [feature_labels.get(feature, feature) for feature in features]
    if len(readable_features) > 1:
        feature_text = ", ".join(readable_features[:-1]) + f" and {readable_features[-1]}"
    else:
        feature_text = readable_features[0] if readable_features else "the selected indicators"
    values = {
        "internal_point": float(engine["H11"].value),
        "release_mode": str(engine["H20"].value),
        "complete_months": int(engine["H18"].value),
        "total_months": 3,
        "pmi_months_available": int(engine["H19"].value),
        "latest_actual": float(summary["A7"].value),
        "latest_period": str(summary["B7"].value),
        "pmi": float(summary["C7"].value),
        "model": str(summary["E7"].value),
        "rmse": float(summary["F7"].value),
        "mae": float(summary["G7"].value),
        "pmi_treatment": str(summary["H7"].value),
        "features": features,
        "feature_text": feature_text,
        "required_inputs": str(engine["H21"].value),
        "quarters": [],
        "pmi_monthly": [],
    }
    for row in monthly.iter_rows(min_row=2, values_only=True):
        observation_date = row[0]
        pmi_value = row[11] if len(row) > 11 else None
        if observation_date is not None and isinstance(pmi_value, (int, float)):
            values["pmi_monthly"].append({"date": observation_date, "value": float(pmi_value)})
    for row in forecast.iter_rows(min_row=4, values_only=True):
        if not row[0]:
            continue
        status = str(row[8] or "")
        actual = float(row[1]) if isinstance(row[1], (int, float)) else None
        point = float(row[5]) if isinstance(row[5], (int, float)) else None
        lower = float(row[6]) if isinstance(row[6], (int, float)) else None
        upper = float(row[7]) if isinstance(row[7], (int, float)) else None
        values["quarters"].append({"period": str(row[0]), "actual": actual, "point": point, "lower": lower, "upper": upper, "status": status})
        if actual is None and status.lower() != "official actual":
            values["target_period"] = str(row[0])
            values["point"] = point
            values["lower"] = lower
            values["upper"] = upper
    wb.close()
    if "target_period" not in values:
        raise RuntimeError("No nowcast row was found in the Quarterly Forecast sheet.")
    return values


def create_gdp_chart(data: dict, chart_path: Path) -> None:
    recent = [r for r in data["quarters"] if r["actual"] is not None][-10:]
    target = next(r for r in data["quarters"] if r["period"] == data["target_period"])
    labels = [quarter_label(r["period"]) for r in recent] + [quarter_label(target["period"])]
    actual = [r["actual"] for r in recent]
    width, height = 1548, 920
    left, right, top, bottom = 150, 55, 105, 190
    plot_left, plot_right = left, width - right
    plot_top, plot_bottom = top, height - bottom
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_dir = Path(r"C:\Windows\Fonts")

    def font(size: int, bold: bool = False):
        preferred = font_dir / ("calibrib.ttf" if bold else "calibri.ttf")
        fallback = font_dir / ("arialbd.ttf" if bold else "arial.ttf")
        return ImageFont.truetype(str(preferred if preferred.exists() else fallback), size)

    label_font = font(27)
    small_font = font(25)
    legend_font = font(27)
    bold_font = font(28, bold=True)
    release_values = [v for v in (data.get("point"), data.get("lower"), data.get("upper")) if v is not None]
    y_min = max(0.0, min(release_values + actual) - 0.7)
    y_max = max(release_values + actual) + 0.7
    y_min = (int(y_min * 2) / 2.0)
    y_max = ((int(y_max * 2) + 1) / 2.0)

    def y_px(value: float) -> float:
        return plot_bottom - (value - y_min) / (y_max - y_min) * (plot_bottom - plot_top)

    x_positions = [plot_left + i * (plot_right - plot_left) / (len(labels) - 1) for i in range(len(labels))]
    tick = y_min
    while tick <= y_max + 1e-9:
        y = y_px(tick)
        draw.line((plot_left, y, plot_right, y), fill="#D9E2F3", width=2)
        txt = f"{tick:.1f}"
        bbox = draw.textbbox((0, 0), txt, font=small_font)
        draw.text((plot_left - 18 - (bbox[2] - bbox[0]), y - (bbox[3] - bbox[1]) / 2), txt, fill="#595959", font=small_font)
        tick += 0.5

    draw.line((plot_left, plot_top, plot_left, plot_bottom), fill="#A6A6A6", width=3)
    draw.line((plot_left, plot_bottom, plot_right, plot_bottom), fill="#A6A6A6", width=3)

    actual_points = [(x_positions[i], y_px(v)) for i, v in enumerate(actual)]
    draw.line(actual_points, fill="#1F4E78", width=7, joint="curve")
    for x, y in actual_points:
        draw.ellipse((x - 8, y - 8, x + 8, y + 8), fill="white", outline="#1F4E78", width=5)

    xf = x_positions[len(actual)]
    if data["release_mode"] == "Indicative range":
        low_y, high_y = y_px(data["lower"]), y_px(data["upper"])
        draw.line((xf, high_y, xf, low_y), fill="#C00000", width=12)
        draw.line((xf - 25, high_y, xf + 25, high_y), fill="#C00000", width=8)
        draw.line((xf - 25, low_y, xf + 25, low_y), fill="#C00000", width=8)
        yf = (low_y + high_y) / 2
    else:
        forecast_points = [(x_positions[len(actual) - 1], y_px(actual[-1])), (xf, y_px(data["point"]))]
        draw.line(forecast_points, fill="#C00000", width=7)
        xf, yf = forecast_points[-1]
        draw.polygon([(xf, yf - 12), (xf + 12, yf), (xf, yf + 12), (xf - 12, yf)], fill="#C00000")

    latest_text = f"{data['latest_actual']:.2f}%"
    latest_box = draw.textbbox((0, 0), latest_text, font=bold_font)
    draw.text((actual_points[-1][0] - (latest_box[2] - latest_box[0]) - 10, actual_points[-1][1] + 18), latest_text, fill="#1F4E78", font=bold_font)
    if data["release_mode"] == "Indicative range":
        forecast_text = f"{data['lower']:.2f}%–{data['upper']:.2f}%"
    else:
        forecast_text = f"{data['point']:.2f}%"
    draw.text((xf - 230, yf - 70), forecast_text, fill="#C00000", font=bold_font)

    for x, label in zip(x_positions, labels):
        bbox = draw.textbbox((0, 0), label, font=label_font)
        draw.text((x - (bbox[2] - bbox[0]) / 2, plot_bottom + 20), label, fill="#404040", font=label_font)

    # Legend.
    legend_y = 28
    draw.line((plot_left, legend_y + 14, plot_left + 52, legend_y + 14), fill="#1F4E78", width=7)
    draw.text((plot_left + 65, legend_y), "NBS actual", fill="#404040", font=legend_font)
    lx2 = plot_left + 300
    if data["release_mode"] == "Indicative range":
        draw.line((lx2 + 26, legend_y - 2, lx2 + 26, legend_y + 31), fill="#C00000", width=10)
        draw.line((lx2 + 10, legend_y - 2, lx2 + 42, legend_y - 2), fill="#C00000", width=6)
        draw.line((lx2 + 10, legend_y + 31, lx2 + 42, legend_y + 31), fill="#C00000", width=6)
        draw.text((lx2 + 65, legend_y), "CAPE indicative range", fill="#404040", font=legend_font)
    else:
        draw.line((lx2, legend_y + 14, lx2 + 52, legend_y + 14), fill="#C00000", width=7)
        draw.text((lx2 + 65, legend_y), "CAPE point forecast", fill="#404040", font=legend_font)

    # Rotated y-axis title.
    y_title = "Real GDP growth (%, year-on-year)"
    title_box = Image.new("RGBA", (800, 70), (255, 255, 255, 0))
    title_draw = ImageDraw.Draw(title_box)
    title_draw.text((0, 0), y_title, fill="#404040", font=label_font)
    title_box = title_box.crop(title_box.getbbox()).rotate(90, expand=True)
    image.paste(title_box, (25, int((plot_top + plot_bottom - title_box.height) / 2)), title_box)

    if data["release_mode"] == "Indicative range":
        note = f"Actual through {quarter_label(data['latest_period'])}; indicative range shown while complete monthly inputs cover {data['complete_months']} of 3 months."
    else:
        note = f"Actual through {quarter_label(data['latest_period'])}; {quarter_label(data['target_period'])} is the CAPE point forecast."
    draw.text((plot_left, height - 80), note, fill="#595959", font=small_font)
    chart_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(chart_path, format="PNG", optimize=True)


def create_pmi_chart(data: dict, chart_path: Path) -> None:
    series = data["pmi_monthly"][-14:]
    if len(series) < 2:
        raise RuntimeError("At least two monthly PMI observations are required for Figure 2.")

    width, height = 1548, 920
    left, right, top, bottom = 145, 55, 105, 190
    plot_left, plot_right = left, width - right
    plot_top, plot_bottom = top, height - bottom
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_dir = Path(r"C:\Windows\Fonts")

    def font(size: int, bold: bool = False):
        preferred = font_dir / ("calibrib.ttf" if bold else "calibri.ttf")
        fallback = font_dir / ("arialbd.ttf" if bold else "arial.ttf")
        return ImageFont.truetype(str(preferred if preferred.exists() else fallback), size)

    label_font = font(24)
    small_font = font(25)
    legend_font = font(27)
    bold_font = font(29, bold=True)
    values = [item["value"] for item in series]
    labels = [item["date"].strftime("%b-%y") for item in series]
    y_min = min(48.0, math.floor(min(values) - 0.5))
    y_max = max(58.0, math.ceil(max(values) + 0.5))

    def y_px(value: float) -> float:
        return plot_bottom - (value - y_min) / (y_max - y_min) * (plot_bottom - plot_top)

    x_positions = [plot_left + i * (plot_right - plot_left) / (len(series) - 1) for i in range(len(series))]
    tick = math.ceil(y_min / 2) * 2
    while tick <= y_max + 1e-9:
        y = y_px(tick)
        draw.line((plot_left, y, plot_right, y), fill="#D9E2F3", width=2)
        txt = f"{tick:.0f}"
        bbox = draw.textbbox((0, 0), txt, font=small_font)
        draw.text((plot_left - 18 - (bbox[2] - bbox[0]), y - (bbox[3] - bbox[1]) / 2), txt, fill="#595959", font=small_font)
        tick += 2

    threshold_y = y_px(50.0)
    dash = 20
    x = plot_left
    while x < plot_right:
        draw.line((x, threshold_y, min(x + dash, plot_right), threshold_y), fill="#C00000", width=4)
        x += dash * 2

    draw.line((plot_left, plot_top, plot_left, plot_bottom), fill="#A6A6A6", width=3)
    draw.line((plot_left, plot_bottom, plot_right, plot_bottom), fill="#A6A6A6", width=3)

    points = [(x_positions[i], y_px(value)) for i, value in enumerate(values)]
    draw.line(points, fill="#1F4E78", width=7, joint="curve")
    for x, y in points[:-1]:
        draw.ellipse((x - 7, y - 7, x + 7, y + 7), fill="white", outline="#1F4E78", width=4)
    latest_x, latest_y = points[-1]
    draw.ellipse((latest_x - 11, latest_y - 11, latest_x + 11, latest_y + 11), fill="#D9A441", outline="#1F4E78", width=4)
    latest_text = f"{values[-1]:.1f}"
    latest_box = draw.textbbox((0, 0), latest_text, font=bold_font)
    draw.text((latest_x - (latest_box[2] - latest_box[0]) - 12, latest_y - 55), latest_text, fill="#1F4E78", font=bold_font)

    for x, label in zip(x_positions, labels):
        bbox = draw.textbbox((0, 0), label, font=label_font)
        draw.text((x - (bbox[2] - bbox[0]) / 2, plot_bottom + 20), label, fill="#404040", font=label_font)

    legend_y = 28
    draw.line((plot_left, legend_y + 14, plot_left + 52, legend_y + 14), fill="#1F4E78", width=7)
    draw.text((plot_left + 65, legend_y), "Nigeria PMI", fill="#404040", font=legend_font)
    lx2 = plot_left + 300
    draw.line((lx2, legend_y + 14, lx2 + 52, legend_y + 14), fill="#C00000", width=4)
    draw.text((lx2 + 65, legend_y), "50-point threshold", fill="#404040", font=legend_font)

    y_title = "PMI index"
    title_box = Image.new("RGBA", (300, 70), (255, 255, 255, 0))
    title_draw = ImageDraw.Draw(title_box)
    title_draw.text((0, 0), y_title, fill="#404040", font=label_font)
    title_box = title_box.crop(title_box.getbbox()).rotate(90, expand=True)
    image.paste(title_box, (30, int((plot_top + plot_bottom - title_box.height) / 2)), title_box)

    change = values[-1] - values[-2]
    draw.text((plot_left, 65), f"Latest monthly change: {change:+.1f} points ({labels[-2]} to {labels[-1]})", fill="#0F6B67", font=small_font)
    px, py = points[-2]
    draw.text((px - 45, py + 18), f"{values[-2]:.1f}", fill="#1F4E78", font=small_font)
    note = "Above 50: improving conditions. Below 50: deteriorating conditions. Values are index levels."
    draw.text((plot_left, height - 80), note, fill="#595959", font=small_font)
    chart_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(chart_path, format="PNG", optimize=True)


def replace_paragraph_text(paragraph, text: str) -> None:
    template_rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            template_rpr = copy.deepcopy(run._r.rPr)
            break
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run(text)
    if template_rpr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, template_rpr)


def replace_caption_keep_drawing(paragraph, text: str) -> None:
    template_rpr = None
    for run in paragraph.runs:
        has_drawing = bool(run._r.xpath(".//*[local-name()='drawing']"))
        if not has_drawing and run._r.rPr is not None and template_rpr is None:
            template_rpr = copy.deepcopy(run._r.rPr)
        if not has_drawing:
            paragraph._p.remove(run._r)
    run = paragraph.add_run(text)
    if template_rpr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, template_rpr)


def replace_source_paragraph(paragraph) -> None:
    lead_rpr = copy.deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    body_rpr = copy.deepcopy(paragraph.runs[1]._r.rPr) if len(paragraph.runs) > 1 and paragraph.runs[1]._r.rPr is not None else lead_rpr
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    lead = paragraph.add_run("Source: ")
    body = paragraph.add_run("National Bureau of Statistics (NBS) and CAPE Economic Research and Consulting")
    for run, rpr in ((lead, lead_rpr), (body, body_rpr)):
        if rpr is not None:
            if run._r.rPr is not None:
                run._r.remove(run._r.rPr)
            run._r.insert(0, copy.deepcopy(rpr))


def find_paragraph(doc: Document, startswith: str):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise RuntimeError(f"Newsletter paragraph not found: {startswith}")


def find_paragraph_any(doc: Document, prefixes: tuple[str, ...]):
    for p in doc.paragraphs:
        text = p.text.strip()
        if any(text.startswith(prefix) for prefix in prefixes):
            return p
    raise RuntimeError(f"Newsletter paragraph not found for any prefix: {prefixes}")


def find_next_drawing_paragraph(doc: Document, paragraph):
    found = False
    for candidate in doc.paragraphs:
        if candidate._p is paragraph._p:
            found = True
            continue
        if found and candidate._p.xpath(".//*[local-name()='blip']"):
            return candidate
    raise RuntimeError("A drawing paragraph was not found after the figure caption.")


def replace_figure_image(doc: Document, paragraph, image_bytes: bytes) -> None:
    blips = paragraph._p.xpath(".//*[local-name()='blip']")
    if not blips:
        raise RuntimeError("Figure image relationship was not found.")
    rid = blips[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
    image_part = doc.part.rels[rid].target_part
    image_part._blob = image_bytes
    # Replacement plots include their own labels; do not inherit template crops.
    for crop in paragraph._p.xpath(".//*[local-name()='srcRect']"):
        crop.getparent().remove(crop)


def apply_plot_overrides(doc: Document, plots_folder: Path | None) -> list[str]:
    if plots_folder is None or not plots_folder.exists():
        return []
    captions_path = plots_folder / "captions.json"
    captions = json.loads(captions_path.read_text(encoding="utf-8")) if captions_path.exists() else {}
    applied = []
    for figure_number in range(1, 8):
        candidates = [
            plots_folder / f"Figure {figure_number}.png",
            plots_folder / f"Figure {figure_number}.jpg",
            plots_folder / f"Figure {figure_number}.jpeg",
        ]
        image_path = next((candidate for candidate in candidates if candidate.exists()), None)
        caption_text = captions.get(str(figure_number))
        if image_path is None and caption_text is None:
            continue
        caption_p = find_paragraph(doc, f"Figure {figure_number}:")
        if caption_text:
            replacement = str(caption_text).strip()
            if not replacement.startswith(f"Figure {figure_number}:"):
                replacement = f"Figure {figure_number}: {replacement}"
            replace_caption_keep_drawing(caption_p, replacement)
        if image_path is not None:
            drawing_p = caption_p if caption_p._p.xpath(".//*[local-name()='blip']") else find_next_drawing_paragraph(doc, caption_p)
            replace_figure_image(doc, drawing_p, image_path.read_bytes())
            applied.append(image_path.name)
    return applied


def publish(
    template: Path,
    workbook: Path,
    output: Path,
    chart: Path,
    pmi_chart: Path,
    report_month: str,
    plots_folder: Path | None,
) -> dict:
    values = read_nowcast(workbook)
    create_gdp_chart(values, chart)
    create_pmi_chart(values, pmi_chart)
    doc = Document(template)
    report_date = datetime.strptime(report_month, "%Y-%m")
    report_month_name = report_date.strftime("%B")
    report_year = report_date.year
    report_title = f"Cape Economic Performance and Prospect Bulletin – {report_month_name} {report_year}"
    replace_paragraph_text(
        find_paragraph_any(
            doc,
            (
                "Cape Economic Performance and Prospect Bulletin",
                "CAPE Economic Performance and Prospect Bulletin",
            ),
        ),
        report_title,
    )

    pmi_series = values["pmi_monthly"]
    pmi_latest = pmi_series[-1]
    pmi_previous = pmi_series[-2]
    consecutive_expansion = 0
    for observation in reversed(pmi_series):
        if observation["value"] > 50.0:
            consecutive_expansion += 1
        else:
            break
    if pmi_latest["value"] > pmi_previous["value"]:
        pmi_change = (
            f"rose to {pmi_latest['value']:.1f} points in {pmi_latest['date'].strftime('%B %Y')} from "
            f"{pmi_previous['value']:.1f} points in {pmi_previous['date'].strftime('%B')}"
        )
    elif pmi_latest["value"] < pmi_previous["value"]:
        pmi_change = (
            f"eased to {pmi_latest['value']:.1f} points in {pmi_latest['date'].strftime('%B %Y')} from "
            f"{pmi_previous['value']:.1f} points in {pmi_previous['date'].strftime('%B')}"
        )
    else:
        pmi_change = (
            f"was unchanged at {pmi_latest['value']:.1f} points in {pmi_latest['date'].strftime('%B %Y')}"
        )
    pmi_status = (
        f"This marks {consecutive_expansion} consecutive months above the 50.0 threshold and signals continued "
        f"private-sector expansion."
        if pmi_latest["value"] > 50.0
        else "The reading is below the 50.0 threshold and signals a deterioration in private-sector conditions."
    )
    pmi_narrative = (
        f"The Stanbic IBTC Nigeria Purchasing Managers’ Index (PMI) {pmi_change}. {pmi_status} "
        f"The improvement points to firmer business activity, although high financing, energy, transport and other "
        f"operating costs remain important constraints."
    )
    replace_paragraph_text(
        find_paragraph_any(
            doc,
            (
                "The Stanbic IBTC Purchasing Managers’ Index",
                "The Stanbic IBTC Nigeria Purchasing Managers’ Index",
                "Nigeria’s Purchasing Managers’ Index",
            ),
        ),
        pmi_narrative,
    )

    pmi_caption_p = find_paragraph(doc, "Figure 2:")
    pmi_caption = (
        f"Figure 2: Trend of Nigeria’s PMI "
        f"({pmi_series[-14]['date'].strftime('%b %Y')}–{pmi_latest['date'].strftime('%b %Y')})"
    )
    replace_caption_keep_drawing(pmi_caption_p, pmi_caption)
    replace_figure_image(doc, find_next_drawing_paragraph(doc, pmi_caption_p), pmi_chart.read_bytes())

    pmi_outlook = (
        f"The latest PMI readings point to a cautious improvement in near-term private-sector activity. The index "
        f"moved from {pmi_previous['value']:.1f} in {pmi_previous['date'].strftime('%B')} to "
        f"{pmi_latest['value']:.1f} in {pmi_latest['date'].strftime('%B')}, remaining above the expansion threshold. "
        f"Sustaining this momentum will require reliable power and transport, access to working capital and continued "
        f"stability in the foreign-exchange market."
    )
    replace_paragraph_text(
        find_paragraph_any(
            doc,
            (
                "Forward-looking indicators remain cautiously positive",
                "The latest PMI readings point to a cautious improvement",
                "The improvement in business activity",
            ),
        ),
        pmi_outlook,
    )

    target_label = quarter_label(values["target_period"])
    latest_label = quarter_label(values["latest_period"])
    if values["release_mode"] == "Indicative range":
        highlight = (
            f"Nigeria’s real GDP growth strengthened to {values['latest_actual']:.2f} per cent in {latest_label}. "
            f"CAPE places {target_label} growth within an indicative range of {values['lower']:.2f}–{values['upper']:.2f} "
            f"per cent. CAPE continues to assess the strength and breadth of the recovery."
        )
    else:
        direction = "slightly below" if values["point"] < values["latest_actual"] else "slightly above"
        highlight = (
            f"Nigeria’s real GDP growth strengthened to {values['latest_actual']:.2f} per cent in {latest_label}; "
            f"CAPE forecasts {target_label} growth at "
            f"{values['point']:.2f} per cent, {direction} the latest official outturn."
        )
    replace_paragraph_text(find_paragraph(doc, "Nigeria’s real GDP growth strengthened"), highlight)

    caption_p = find_paragraph(doc, "Figure 3:")
    release_label = "Indicative Range" if values["release_mode"] == "Indicative range" else "Point Forecast"
    caption = f"Figure 3: Nigeria Real GDP Growth and CAPE {target_label} {release_label} (Q1 2024–{target_label}) in percent"
    replace_caption_keep_drawing(caption_p, caption)
    replace_figure_image(doc, caption_p, chart.read_bytes())

    source_p = find_paragraph(doc, "Source: National Bureau of Statistics (NBS)")
    replace_source_paragraph(source_p)

    if values["release_mode"] == "Indicative range":
        outlook = (
            f"CAPE places {target_label} real GDP growth in an indicative {values['lower']:.2f}–{values['upper']:.2f} per cent range. "
            f"High borrowing costs, energy expenses and insecurity remain risks to the outlook. "
            f"We continue to monitor economic developments and will refine the outlook as the quarter progresses."
        )
    else:
        outlook = (
            f"Nigeria’s near-term output outlook remains subject to domestic and external risks. "
            f"CAPE forecasts real GDP growth at {values['point']:.2f} per cent year-on-year in {target_label}, "
            f"{direction} the {values['latest_actual']:.2f} per cent recorded in {latest_label}. "
            f"High financing costs, energy and logistics expenses, insecurity and weak household demand remain "
            f"downside risks."
        )
    replace_paragraph_text(
        find_paragraph_any(
            doc,
            (
                "Nigeria’s near-term output outlook",
                f"CAPE places {target_label} real GDP growth",
                "CAPE maintains an indicative range",
            ),
        ),
        outlook,
    )

    conclusion_p = find_paragraph(doc, "Nigeria enters ")
    if values["release_mode"] == "Indicative range":
        growth_sentence = (
            f"CAPE places {target_label} growth in an indicative {values['lower']:.2f}–{values['upper']:.2f} per cent "
            f"range, subject to further assessment as the quarter progresses."
        )
    else:
        growth_sentence = (
            f"CAPE forecasts {values['point']:.2f} per cent growth in {target_label}."
        )
    conclusion = re.sub(r"^Nigeria enters [A-Za-z]+", f"Nigeria enters {report_month_name}", conclusion_p.text.strip())
    growth_start = conclusion.find("Real GDP growth")
    continuation_markers = [" Stronger federation", " CAPE ERC expects", " The outlook will depend"]
    continuation_positions = [conclusion.find(marker, growth_start) for marker in continuation_markers if conclusion.find(marker, growth_start) >= 0]
    growth_end = min(continuation_positions) if continuation_positions else -1
    refreshed_growth = f"Real GDP growth accelerated to {values['latest_actual']:.2f} per cent in {latest_label}. {growth_sentence}"
    if growth_start >= 0 and growth_end > growth_start:
        conclusion = conclusion[:growth_start] + refreshed_growth + conclusion[growth_end:]
    replace_paragraph_text(conclusion_p, conclusion)

    plot_overrides = apply_plot_overrides(doc, plots_folder)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return {
        **values,
        "template": str(template),
        "workbook": str(workbook),
        "output": str(output),
        "chart": str(chart),
        "pmi_chart": str(pmi_chart),
        "report_month": report_month,
        "plot_overrides": plot_overrides,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Publish the Nigeria GDP nowcast into the CAPE economic newsletter.")
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--chart", type=Path)
    parser.add_argument("--pmi-chart", type=Path)
    parser.add_argument("--report-month", default=datetime.now().strftime("%Y-%m"))
    parser.add_argument("--plots-folder", type=Path)
    args = parser.parse_args()
    chart = args.chart or args.workbook.parent / "Newsletter Assets" / "Nigeria GDP Nowcast Chart.png"
    pmi_chart = args.pmi_chart or args.workbook.parent / "Newsletter Assets" / "Nigeria PMI Trend Chart.png"
    template = args.template.resolve()
    output = args.output.resolve()
    if not template.exists() and output.exists():
        template = output
    plots_folder = args.plots_folder.resolve() if args.plots_folder else None
    result = publish(template, args.workbook.resolve(), output, chart.resolve(), pmi_chart.resolve(), args.report_month, plots_folder)
    print({
        "status": "newsletter published",
        "target": result["target_period"],
        "release_mode": result["release_mode"],
        "point": result["point"],
        "published_range": [result["lower"], result["upper"]],
        "report_month": result["report_month"],
        "plot_overrides": result["plot_overrides"],
        "output": result["output"],
    })


if __name__ == "__main__":
    main()
